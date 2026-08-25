"""Export a ranked shortlist to a spreadsheet you can work through by hand.

This is the manual-apply companion to the automatic run. It searches the
cities you name, scores everything against your profile, and writes the top N
to an .xlsx with the apply link on each row plus columns for tracking what you
sent and where.

The tracking columns matter more than they look. Applications go out through
two portals - Naukri and LinkedIn - and the same job is often listed on both.
Without one sheet recording which one you used, you end up either applying
twice to the same posting or skipping it on both.

Anything already in data/jobs/ledger.json is marked as applied, so the agent's
own applications and your manual ones stay in the same picture.

Two optional cuts turn this from a standing shortlist into a daily digest:
`posted_days` keeps only listings posted inside a window (1 = the last 24
hours) and `new_only` keeps only jobs that have not appeared on an earlier
day's page. Both are applied before the top-N truncation, so what survives is
a full list of what changed rather than the leftovers of a stale one.
"""
from __future__ import annotations

import logging
from datetime import date
from pathlib import Path
from urllib.parse import quote_plus

from . import score as score_mod
from .ledger import Ledger
from .page import age_days as label_age_days, linkedin_posted, load_seen

log = logging.getLogger("screener.export")

from .paths import JOBS_DIR

HEADERS = [
    ("#", 5),
    ("Score", 7),
    ("Title", 42),
    ("Company", 24),
    ("Location", 26),
    ("Experience", 12),
    ("Salary", 16),
    ("Posted", 12),
    ("Matched skills", 40),
    ("Apply on Naukri", 16),
    ("Find on LinkedIn", 16),
    ("Applied?", 10),
    ("Where", 12),
    ("Applied on", 12),
    ("Notes", 30),
]

# Naukri writes the current official names, which are not what people search
# for. A "Gurgaon" search returns listings labelled "Gurugram", so filtering on
# the search term alone silently drops every job in the city you asked for.
CITY_ALIASES = {
    "gurgaon": ("gurgaon", "gurugram"),
    "gurugram": ("gurgaon", "gurugram"),
    "bangalore": ("bangalore", "bengaluru"),
    "bengaluru": ("bangalore", "bengaluru"),
    "bombay": ("bombay", "mumbai"),
    "mumbai": ("bombay", "mumbai"),
    "calcutta": ("calcutta", "kolkata"),
    "kolkata": ("calcutta", "kolkata"),
    "madras": ("madras", "chennai"),
    "chennai": ("madras", "chennai"),
    "noida": ("noida", "greater noida"),
    "trivandrum": ("trivandrum", "thiruvananthapuram"),
}


def city_variants(city: str) -> tuple[str, ...]:
    key = (city or "").lower().strip()
    return CITY_ALIASES.get(key, (key,))


def posted_within(job, days: float | None) -> bool:
    """True when a Naukri listing was posted inside the last `days` days.

    Naukri stamps every search result with a createdDate, so this is normally
    an exact age in hours rather than a reading of "3 Days Ago". The label is
    only a fallback, and a listing carrying neither is dropped: a freshness
    filter that keeps undated rows is not a freshness filter.
    """
    if days is None:
        return True
    age = job.age_days
    if age is None:
        label = label_age_days(job.posted_label)
        age = None if label is None else float(label)
    return age is not None and age <= days


def card_posted_within(card: dict, days: float | None) -> bool:
    """True when a LinkedIn card is inside the last `days` days.

    Undated cards are KEPT here, unlike the Naukri side. LinkedIn applies its
    own f_TPR window server-side before the results are ever rendered, so an
    undated card is one LinkedIn already considers in-window - promoted cards
    show "Promoted" where the date would be. Only a card that states an age
    past the window is dropped.
    """
    if days is None:
        return True
    age = label_age_days(linkedin_posted(card))
    return age is None or age <= days


STATUS_CHOICES = '"To apply,Applied,Shortlisted,Rejected,Not interested"'
PORTAL_CHOICES = '"Naukri,LinkedIn,Company site,Referral"'


def linkedin_search_url(job) -> str:
    """A LinkedIn job search for the same role at the same company.

    LinkedIn has no stable per-posting URL we can derive from a Naukri
    listing, so this is a search rather than a deep link - it lands you on the
    same role if LinkedIn carries it.
    """
    terms = " ".join(part for part in (job.title, job.company) if part)
    return f"https://www.linkedin.com/jobs/search/?keywords={quote_plus(terms)}"


def to_excel(jobs: list, path: Path, ledger: Ledger | None = None,
             linkedin_cards: list[dict] | None = None, config: dict | None = None) -> Path:
    """Write the ranked jobs to an .xlsx. Returns the path written."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    ledger = ledger if ledger is not None else Ledger()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Matches"

    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(bold=True, color="FFFFFF")
    for column, (title, width) in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 28

    link_font = Font(color="0563C1", underline="single")
    strong_fill = PatternFill("solid", fgColor="E2EFDA")   # score >= 72
    applied_fill = PatternFill("solid", fgColor="FFF2CC")

    for index, job in enumerate(jobs, start=1):
        row = index + 1
        status = ledger.status(job.job_id)
        already = status == "applied"
        matched = ", ".join(getattr(job, "matched_skills", [])[:8])

        values = [
            index,
            getattr(job, "score", None),
            job.title,
            job.company,
            job.location or "",
            job.experience_label or "",
            job.salary_label or "",
            job.posted_label or "",
            matched,
            None,  # Naukri link
            None,  # LinkedIn link
            "Applied" if already else "To apply",
            "Naukri" if already else "",
            date.today().isoformat() if already else "",
            "applied by the agent" if already else "",
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=column in (3, 4, 5, 9, 15))

        naukri_cell = sheet.cell(row=row, column=10, value="Apply")
        naukri_cell.hyperlink = job.url
        naukri_cell.font = link_font

        linkedin_cell = sheet.cell(row=row, column=11, value="Search")
        linkedin_cell.hyperlink = linkedin_search_url(job)
        linkedin_cell.font = link_font

        if already:
            for column in range(1, len(HEADERS) + 1):
                sheet.cell(row=row, column=column).fill = applied_fill
        elif (getattr(job, "score", 0) or 0) >= 72:
            sheet.cell(row=row, column=2).fill = strong_fill

    last_row = len(jobs) + 1

    # Dropdowns, so the tracking columns stay consistent enough to filter on.
    status_validation = DataValidation(type="list", formula1=STATUS_CHOICES, allow_blank=True)
    portal_validation = DataValidation(type="list", formula1=PORTAL_CHOICES, allow_blank=True)
    sheet.add_data_validation(status_validation)
    sheet.add_data_validation(portal_validation)
    status_validation.add(f"L2:L{max(last_row, 2)}")
    portal_validation.add(f"M2:M{max(last_row, 2)}")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(last_row, 2)}"

    if linkedin_cards:
        add_linkedin_sheet(workbook, linkedin_cards, config or {}, ledger)
    _add_notes_sheet(workbook, jobs)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    log.info("Wrote %d job(s) to %s", len(jobs), path)
    return path


def _add_notes_sheet(workbook, jobs: list) -> None:
    """A second tab explaining the score, so the ranking is auditable."""
    from openpyxl.styles import Font

    sheet = workbook.create_sheet("Why these ranked")
    sheet.column_dimensions["A"].width = 5
    sheet.column_dimensions["B"].width = 42
    sheet.column_dimensions["C"].width = 24
    sheet.column_dimensions["D"].width = 90

    for column, title in enumerate(["#", "Title", "Company", "Score breakdown"], start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.font = Font(bold=True)

    for index, job in enumerate(jobs, start=1):
        sheet.cell(row=index + 1, column=1, value=index)
        sheet.cell(row=index + 1, column=2, value=job.title)
        sheet.cell(row=index + 1, column=3, value=job.company)
        sheet.cell(row=index + 1, column=4, value=score_mod.explain(job))
    sheet.freeze_panes = "A2"


def run(locations: list[str], top: int = 30, headless: bool = False,
        include_applied: bool = False, include_linkedin: bool = False,
        worldwide: bool = False, posted_days: float | None = None,
        new_only: bool = False) -> tuple[Path, list, list]:
    """Search the given cities, rank, and write the spreadsheet.

    `posted_days` restricts the run to listings posted inside that window
    (1 = the last 24 hours), asked for at each board's own search filter and
    then enforced again on the results. `new_only` drops anything that has
    already appeared on an earlier day's tracker page, so a daily scan reports
    what changed since yesterday instead of re-listing the same postings.
    """
    from playwright.sync_api import sync_playwright

    from .session import DEFAULT_STATE, open_profile
    from . import config as config_mod, search

    profile = config_mod.load_profile()
    config = config_mod.load(profile=profile)

    keywords, seen = [], set()
    for entry in config.get("searches") or []:
        keyword = (entry or {}).get("keyword")
        if keyword and keyword.lower() not in seen:
            seen.add(keyword.lower())
            keywords.append(keyword)

    config = dict(config)
    if worldwide:
        # Naukri is an India-only board, so the widest it goes is a national
        # search with no city. "Worldwide" only really means anything on the
        # LinkedIn side.
        config["searches"] = [{"keyword": keyword, "location": None} for keyword in keywords]
    else:
        config["searches"] = [
            {"keyword": keyword, "location": location}
            for location in locations
            for keyword in keywords
        ]
    # Score every requested city equally. jobs.yaml lists only Pune as
    # preferred, which costs every Ahmedabad and Gurugram listing 8 points and
    # buries them - on a list you asked to span three cities, that is just a
    # Pune list with extra steps.
    config["preferred_locations"] = ["Remote"] if worldwide else list(locations) + ["Remote"]
    # Read by search.gather, which turns it into Naukri's own jobAge facet.
    config["posted_within_days"] = posted_days
    log.info("Searching %d keyword(s) across %s", len(keywords), ", ".join(locations))
    if posted_days:
        log.info("Freshness filter: posted in the last %g day(s)", posted_days)

    # Jobs that appeared on an earlier day's page. Today's own entries are not
    # excluded, so a second run on the same day reports the same list rather
    # than an empty one.
    today = date.today().isoformat()
    already_listed: set[str] = set()
    if new_only:
        already_listed = {
            job_id for job_id, first in load_seen().items() if first != today
        }
        log.info("New-only: %d job(s) listed on an earlier day will be skipped",
                 len(already_listed))

    ledger = Ledger()
    with sync_playwright() as p:
        browser, _ctx, page = open_profile(p, DEFAULT_STATE, headless=headless)
        try:
            jobs = search.gather(page, config)
        finally:
            browser.close()

    wanted = [variant for loc in locations for variant in city_variants(loc)]
    kept = []
    stale = repeats = 0
    for job in jobs:
        score_mod.score(job, config)
        if job.score <= 0:
            continue
        # Naukri happily returns Bengaluru jobs for a Pune search, so filter on
        # the location the listing actually states.
        text = (job.location or "").lower()
        if not worldwide and not any(city in text for city in wanted) and "remote" not in text:
            continue
        if not include_applied and ledger.status(job.job_id) == "applied":
            continue
        # Both cuts happen before the top-N truncation, so the shortlist fills
        # with fresh, unseen jobs rather than being trimmed down to a handful.
        if not posted_within(job, posted_days):
            stale += 1
            continue
        if f"naukri:{job.job_id}" in already_listed:
            repeats += 1
            continue
        kept.append(job)

    if posted_days:
        log.info("Naukri: dropped %d listing(s) older than %g day(s)", stale, posted_days)
    if new_only:
        log.info("Naukri: dropped %d listing(s) already seen on an earlier day", repeats)

    # Remote first when asked for, then by score.
    if worldwide:
        kept.sort(key=lambda j: ("remote" not in (j.location or "").lower(), -j.score))
    else:
        kept.sort(key=lambda j: j.score, reverse=True)
    kept = kept[:top]

    cards = []
    if include_linkedin:
        try:
            cards = gather_linkedin(config, locations, top=top, headless=headless,
                                    worldwide=worldwide,
                                    posted_days=int(posted_days) if posted_days else 30,
                                    exclude_ids=already_listed)
            log.info("LinkedIn: %d card(s) after filtering", len(cards))
        except Exception as exc:
            # A LinkedIn failure must not cost you the Naukri sheet.
            log.warning("LinkedIn search skipped: %s", str(exc)[:200])

    path = JOBS_DIR / f"job-matches-{date.today().isoformat()}.xlsx"
    to_excel(kept, path, ledger, linkedin_cards=cards, config=config)

    # Persist the raw results so the HTML page can be rebuilt without
    # re-running the scan.
    import json
    results = {
        "generated": date.today().isoformat(),
        "locations": locations,
        "worldwide": worldwide,
        "posted_days": posted_days,
        "new_only": new_only,
        "naukri": [
            dict(job.to_dict(), score=job.score,
                 matched_skills=getattr(job, "matched_skills", []))
            for job in kept
        ],
        "linkedin": cards,
    }
    (JOBS_DIR / f"results-{date.today().isoformat()}.json").write_text(
        json.dumps(results, indent=2, ensure_ascii=False), encoding="utf-8")

    from . import page as page_mod
    results["_page"] = str(page_mod.build(results))

    return path, kept, cards


def summarise(path: Path, jobs: list, locations: list[str], cards: list | None = None,
              posted_days: float | None = None, new_only: bool = False) -> str:
    strong = sum(1 for j in jobs if (getattr(j, "score", 0) or 0) >= 72)
    lines = [
        "",
        f"  {len(jobs)} job(s) across {', '.join(locations)}",
        f"  {strong} scoring 72+ (the auto-apply bar)",
    ]
    if posted_days:
        window = "the last 24 hours" if posted_days == 1 else f"the last {posted_days:g} days"
        lines.append(f"  Posted in {window} only")
    if new_only:
        lines.append("  First seen today only - anything listed on an earlier day is excluded")
    if posted_days and not jobs:
        lines.append("  Nothing new in that window. That is a normal result for a "
                     "24-hour filter, not a broken scan.")
    lines += ["", "  Top 10:"]
    for index, job in enumerate(jobs[:10], start=1):
        lines.append(
            f"   {index:2}. {job.score:5.1f}  {job.title[:42]:44} {job.company[:22]:24} {job.location or ''}"
        )
    lines += ["", f"  Spreadsheet: {path}", ""]
    return "\n".join(lines)


# --------------------------------------------------------------- LinkedIn tab

LINKEDIN_HEADERS = [
    ("#", 5),
    ("Title match", 11),
    ("Title", 46),
    ("Company", 26),
    ("Location", 34),
    ("Salary", 16),
    ("Easy Apply", 11),
    ("Promoted", 10),
    ("Open on LinkedIn", 16),
    ("Applied?", 10),
    ("Where", 12),
    ("Applied on", 12),
    ("Notes", 26),
]


class _TitleOnly:
    """Minimal stand-in so the title scorer can read a LinkedIn card."""

    def __init__(self, title: str):
        self.title = title or ""


def title_match(card: dict, config: dict) -> int:
    """How well a LinkedIn card's title matches what you are looking for, 0-100.

    Deliberately NOT the 0-100 used for Naukri. A LinkedIn card carries no
    skills list and no experience range - the two components worth 60 of the
    Naukri score - so putting both on one scale would be fake precision. This
    is title fit alone, and the column is labelled as such.
    """
    return round(score_mod._title_score(_TitleOnly(card.get("title")), config) / 25 * 100)


def salary_of(card: dict) -> str:
    for item in card.get("metadata") or []:
        if any(mark in item for mark in ("/yr", "/hr", "₹", "$", "LPA", "INR")):
            return item
    return ""


def add_linkedin_sheet(workbook, cards: list[dict], config: dict, ledger: Ledger) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    sheet = workbook.create_sheet("LinkedIn")
    header_fill = PatternFill("solid", fgColor="0A66C2")
    header_font = Font(bold=True, color="FFFFFF")
    for column, (title, width) in enumerate(LINKEDIN_HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 28

    link_font = Font(color="0563C1", underline="single")

    for index, card in enumerate(cards, start=1):
        row = index + 1
        values = [
            index,
            card.get("_match"),
            card.get("title"),
            card.get("company"),
            card.get("location"),
            salary_of(card),
            "Yes" if card.get("easy_apply") else "",
            "Yes" if card.get("promoted") else "",
            None,
            "To apply",
            "",
            "",
            card.get("insight") or "",
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=column in (3, 4, 5, 13))

        link_cell = sheet.cell(row=row, column=9, value="Open")
        link_cell.hyperlink = card["url"]
        link_cell.font = link_font

    last_row = max(len(cards) + 1, 2)
    status_validation = DataValidation(type="list", formula1=STATUS_CHOICES, allow_blank=True)
    portal_validation = DataValidation(type="list", formula1=PORTAL_CHOICES, allow_blank=True)
    sheet.add_data_validation(status_validation)
    sheet.add_data_validation(portal_validation)
    status_validation.add(f"J2:J{last_row}")
    portal_validation.add(f"K2:K{last_row}")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(LINKEDIN_HEADERS))}{last_row}"


def gather_linkedin(config: dict, locations: list[str], posted_days: int = 30,
                    top: int = 30, headless: bool = False,
                    worldwide: bool = False,
                    exclude_ids: set[str] | None = None) -> list[dict]:
    """Search LinkedIn for the configured keywords across the given cities.

    `posted_days` goes straight into LinkedIn's own f_TPR window, so the
    freshness cut happens at the source. `exclude_ids` holds "linkedin:<id>"
    keys already listed on an earlier day, dropped before the top-N cut.
    """
    from playwright.sync_api import sync_playwright

    from . import linkedin as linkedin_mod

    keywords, seen = [], set()
    for entry in config.get("searches") or []:
        keyword = (entry or {}).get("keyword")
        if keyword and keyword.lower() not in seen:
            seen.add(keyword.lower())
            keywords.append(keyword)

    cards: dict[str, dict] = {}
    with sync_playwright() as p:
        browser, _ctx, page = linkedin_mod.open_session(p, headless=headless)
        try:
            if worldwide:
                # Remote first, and via LinkedIn's own workplace-type filter
                # rather than the word "remote", which also matches on-site
                # roles whose description happens to mention remote working.
                for keyword in keywords:
                    for card in linkedin_mod.search(page, keyword, None, posted_days,
                                                    pages=2, remote_only=True):
                        cards.setdefault(card["job_id"], card)
                    linkedin_mod.pause()
                for keyword in keywords:
                    for card in linkedin_mod.search(page, keyword, None, posted_days, pages=1):
                        cards.setdefault(card["job_id"], card)
                    linkedin_mod.pause()
            else:
                for location in locations:
                    # LinkedIn wants a region string, not a bare city name.
                    place = location if "," in location else f"{location}, India"
                    for keyword in keywords:
                        for card in linkedin_mod.search(page, keyword, place, posted_days, pages=1):
                            cards.setdefault(card["job_id"], card)
                        linkedin_mod.pause()
        finally:
            browser.close()

    wanted = [variant for loc in locations for variant in city_variants(loc)]
    window = posted_days if posted_days and posted_days < 30 else None
    kept = []
    for card in cards.values():
        text = (card.get("location") or "").lower()
        if not worldwide and not any(city in text for city in wanted) and "remote" not in text:
            continue
        if exclude_ids and f"linkedin:{card.get('job_id')}" in exclude_ids:
            continue
        if not card_posted_within(card, window):
            continue
        card["_match"] = title_match(card, config)
        card["_remote"] = bool(card.get("remote_filtered")) or "remote" in text
        kept.append(card)

    # Remote first, then title fit. Organic before promoted at equal fit - a
    # paid placement is not evidence the role suits you.
    kept.sort(key=lambda c: (not c["_remote"], -c["_match"], bool(c.get("promoted"))))
    return kept[:top]
